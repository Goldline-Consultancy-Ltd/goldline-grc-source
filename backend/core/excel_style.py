"""Goldline GRC house style for the spreadsheets the app generates.

Every export builds its workbook with openpyxl (directly or through pandas), so
the style is applied once, just before openpyxl writes the file: Calibri, a navy
header row with bold white text, light borders, rating fills, a frozen header
row and column widths that fit. Workbooks loaded from a file (for example the
official CyFun and ANSSI forms the app fills in) are left exactly as they are.
"""

from openpyxl.reader.excel import ExcelReader
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.writer.excel import ExcelWriter

from core import report_brand as brand

MAX_STYLED_ROWS = 20000
MIN_WIDTH, MAX_WIDTH = 8, 60

_side = Side(style="thin", color=brand.BORDER)
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
HEADER_FILL = PatternFill("solid", fgColor=brand.NAVY)
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color=brand.WHITE)
BODY_FONT = Font(name="Calibri", size=10, color=brand.TEXT)
RATING_FILLS = {
    word: PatternFill("solid", fgColor=colour)
    for colour, words in brand.RATING_WORDS.items()
    for word in words
}


def _looks_like_header(row) -> bool:
    values = [c.value for c in row if c.value not in (None, "")]
    return bool(values) and all(isinstance(v, str) for v in values)


def style_worksheet(ws):
    if ws.max_row < 1 or ws.max_column < 1:
        return
    last_row = min(ws.max_row, MAX_STYLED_ROWS)
    header = _looks_like_header(ws[1])
    widths = {}
    for row in ws.iter_rows(min_row=1, max_row=last_row):
        for cell in row:
            if cell.__class__.__name__ == "MergedCell":
                continue
            value = cell.value
            is_header = header and cell.row == 1
            cell.font = HEADER_FONT if is_header else BODY_FONT
            cell.border = BORDER
            if is_header:
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                fill = RATING_FILLS.get(str(value).strip().lower()) if isinstance(value, str) else None
                if fill is not None:
                    cell.fill = fill
                long_text = isinstance(value, str) and len(value) > MAX_WIDTH
                cell.alignment = Alignment(vertical="top", wrap_text=long_text)
            if value is not None:
                longest = max((len(part) for part in str(value).splitlines()), default=0)
                widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(MIN_WIDTH, min(MAX_WIDTH, width + 2))
    if header:
        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22


def style_workbook(wb):
    for ws in wb.worksheets:
        style_worksheet(ws)


def apply_excel_house_style():
    if getattr(ExcelWriter, "_goldline_style", False):
        return

    original_read = ExcelReader.read
    original_write = ExcelWriter.write_data

    def read(self, *args, **kwargs):
        result = original_read(self, *args, **kwargs)
        self.wb._goldline_loaded = True
        return result

    def write_data(self, *args, **kwargs):
        wb = self.workbook
        if not (getattr(wb, "_goldline_loaded", False) or getattr(wb, "write_only", False)):
            style_workbook(wb)
        return original_write(self, *args, **kwargs)

    ExcelReader.read = read
    ExcelWriter.write_data = write_data
    ExcelWriter._goldline_style = True
