"""Build the Goldline house-style 16:9 slide template.

Writes core/templates/core/audit_report_template_en.pptx: a cover, contents,
summary (rating table) and findings slide in the house colours, ready to copy
for executive or summary decks. The app has no deck export of its own, so the
slides carry [square-bracket] fields to fill in.

Needs python-pptx (not a backend dependency):
    pip install python-pptx==1.0.2 && python scripts/goldline_build_deck_template.py
"""

import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
from core import report_brand as brand  # noqa: E402

OUT = BACKEND / "core" / "templates" / "core" / "audit_report_template_en.pptx"
LOGO = BACKEND / "core" / "templates" / "core" / "goldline-grc-mark.png"
FONT = "Calibri"
W, H = Mm(338.67), Mm(190.5)  # 16:9
MARGIN = Mm(20)


def rgb(hex_):
    return RGBColor.from_string(hex_)


def text(slide, x, y, w, h, value, size, colour, bold=False, align=PP_ALIGN.LEFT, caps=False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value.upper() if caps else value
    f = r.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = rgb(colour)
    return box


def rule(slide, y, weight_pt, colour=brand.GOLD, x=MARGIN, width=None):
    width = width or (W - 2 * MARGIN)
    shape = slide.shapes.add_connector(1, x, y, x + width, y)
    shape.line.color.rgb = rgb(colour)
    shape.line.width = Pt(weight_pt)
    return shape


def chrome(slide, title):
    """Running header (client | title, gold rule) and the company footer."""
    text(slide, MARGIN, Mm(8), W - 2 * MARGIN, Mm(5), f"[Client name]  |  {title}", 7.5, brand.BODY, align=PP_ALIGN.RIGHT)
    rule(slide, Mm(14), 0.75)
    box = text(slide, MARGIN, H - Mm(12), W - 2 * MARGIN, Mm(6), brand.footer_text() + "  Page ", 7.5, brand.BODY, align=PP_ALIGN.CENTER)
    r = box.text_frame.paragraphs[0].add_run()
    r.font.name, r.font.size = FONT, Pt(9.5)
    r.font.color.rgb = rgb(brand.TEXT)
    fld = r._r
    # Slide-number field.
    new = fld.makeelement(qn("a:fld"), {"id": "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", "type": "slidenum"})
    new.append(fld.find(qn("a:rPr")))
    t = new.makeelement(qn("a:t"), {})
    t.text = "#"
    new.append(t)
    fld.addprevious(new)
    fld.getparent().remove(fld)


def table(slide, x, y, w, rows, header, widths_mm, fills=None):
    shape = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, Mm(8) * (len(rows) + 1))
    tbl = shape.table
    # Plain table: no built-in banding style.
    tblpr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow"):
        tblpr.set(attr, "0")
    style = tblpr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # "No Style, Table Grid"
    for i, wmm in enumerate(widths_mm):
        tbl.columns[i].width = Mm(wmm)
    all_rows = [header] + rows
    for r, values in enumerate(all_rows):
        for c, value in enumerate(values):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = cell.margin_top = cell.margin_bottom = Pt(4)
            fill = brand.NAVY if r == 0 else ((fills or {}).get((r, c)) or brand.WHITE)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = value
            run.font.name, run.font.size, run.font.bold = FONT, Pt(9), r == 0
            run.font.color.rgb = rgb(brand.WHITE if r == 0 else brand.TEXT)
            borders(cell)
    return shape


def borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    for edge in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcpr.findall(qn(edge)):
            tcpr.remove(old)
        ln = tcpr.makeelement(qn(edge), {"w": str(Pt(0.5)), "cmpd": "sng"})
        solid = ln.makeelement(qn("a:solidFill"), {})
        clr = solid.makeelement(qn("a:srgbClr"), {"val": brand.BORDER})
        solid.append(clr)
        ln.append(solid)
        # Line properties go before the cell fill.
        fill = tcpr.find(qn("a:solidFill"))
        if fill is not None:
            fill.addprevious(ln)
        else:
            tcpr.append(ln)


def heading(slide, value, y=Mm(20)):
    text(slide, MARGIN, y, W - 2 * MARGIN, Mm(10), value, 15, brand.NAVY, bold=True)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]
    title = "[Report title]"

    # Cover.
    s = prs.slides.add_slide(blank)
    chrome(s, title)
    s.shapes.add_picture(str(LOGO), MARGIN, Mm(20), height=Mm(20))
    text(s, MARGIN, Mm(48), Mm(180), Mm(6), "Compliance assessment", 10, brand.GOLD, bold=True, caps=True)
    text(s, MARGIN, Mm(56), Mm(180), Mm(12), "[Report title]", 22, brand.NAVY, bold=True)
    text(s, MARGIN, Mm(69), Mm(180), Mm(9), "Mapped to ISO/IEC 27001:2022", 15, brand.NAVY)
    text(s, MARGIN, Mm(80), Mm(180), Mm(7), "Executive summary", 12.5, brand.TEXT, bold=True)
    rule(s, Mm(90), 2, width=Mm(180))
    rows = [
        ["Company assessed", "[Company assessed]"],
        ["Prepared for", "[Prepared for]"],
        ["Engagement", "[Engagement]"],
        ["Report date", "[Report date]"],
        ["Version", "[Version]"],
        ["Assessor", brand.DEFAULT_ASSESSOR],
    ]
    shape = table(s, Mm(210), Mm(30), Mm(108), rows[1:], rows[0], (40, 68))
    # The details table has label cells on cream, not a navy header.
    tbl = shape.table
    for r in range(len(rows)):
        for c in range(2):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(brand.CREAM if c == 0 else brand.WHITE)
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.bold = c == 0
            run.font.color.rgb = rgb(brand.TEXT)
    box = text(s, MARGIN, Mm(100), Mm(180), Mm(20), "Confidential working document.", 9, brand.TEXT, bold=True)
    r = box.text_frame.paragraphs[0].add_run()
    r.text = (
        f" Prepared for [Client name] and intended only for its management and auditors. Do not share it with "
        f"third parties without written agreement from {brand.COMPANY}. Governed by the law of {brand.JURISDICTION}."
    )
    r.font.name, r.font.size = FONT, Pt(9)
    r.font.color.rgb = rgb(brand.BODY)

    # Contents.
    s = prs.slides.add_slide(blank)
    chrome(s, title)
    heading(s, "Contents")
    for i, item in enumerate(("1. Summary", "2. Findings", "3. Action plan", "4. Next steps")):
        text(s, MARGIN, Mm(36 + i * 9), Mm(200), Mm(7), item, 12.5, brand.TEXT, bold=True)

    # Summary with rating table.
    s = prs.slides.add_slide(blank)
    chrome(s, title)
    heading(s, "1. Summary")
    box = text(s, MARGIN, Mm(33), W - 2 * MARGIN, Mm(8), "What this covers:", 10.5, brand.TEXT, bold=True)
    r = box.text_frame.paragraphs[0].add_run()
    r.text = " [one or two sentences on scope and overall position]"
    r.font.name, r.font.size = FONT, Pt(10.5)
    r.font.color.rgb = rgb(brand.BODY)
    labels = ["Met", "Partially met", "Not met", "Not applicable", "Not assessed"]
    keys = ["compliant", "partially_compliant", "non_compliant", "not_applicable", "not_assessed"]
    fills = {(1, i): brand.RATING_FILL[k] for i, k in enumerate(keys)}
    table(s, MARGIN, Mm(46), W - 2 * MARGIN, [["[n]"] * 5], labels, [59.7] * 5, fills)

    # Findings table.
    s = prs.slides.add_slide(blank)
    chrome(s, title)
    heading(s, "2. Findings")
    rows = [
        ["[A.5.15]", "[Requirement]", "[Observation]", "Met"],
        ["[A.8.13]", "[Requirement]", "[Observation]", "Partially met"],
        ["[A.5.24]", "[Requirement]", "[Observation]", "Not met"],
        ["[A.8.8]", "[Requirement]", "[Observation]", "Not assessed"],
    ]
    fills = {(i + 1, 3): brand.RATING_FILL[k] for i, k in enumerate(["compliant", "partially_compliant", "non_compliant", "not_assessed"])}
    table(s, MARGIN, Mm(34), W - 2 * MARGIN, rows, ["ISO 27001 reference", "Requirement", "Observation", "Rating"], [40, 90, 130, 38.7], fills)

    cp = prs.core_properties
    cp.author = cp.last_modified_by = brand.COMPANY
    cp.title = "Goldline report slides"
    return prs


def set_company(path: Path):
    with zipfile.ZipFile(path) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in entries:
            if info.filename == "docProps/app.xml":
                xml = re.sub(r"<Company>.*?</Company>|<Company/>", "", data.decode())
                data = xml.replace("</Properties>", f"<Company>{brand.COMPANY}</Company></Properties>").encode()
            z.writestr(info, data)


if __name__ == "__main__":
    build().save(OUT)
    set_company(OUT)
    print("wrote", OUT)
