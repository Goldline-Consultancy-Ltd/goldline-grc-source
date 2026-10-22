"""Build the Goldline house-style Word audit report templates.

Writes core/templates/core/audit_report_template_en.docx, _fr.docx and the
language-neutral audit_report_template.docx (English). The templates are
docxtpl/Jinja documents filled by `ComplianceAssessmentViewSet.word_report`
with `gen_audit_context` plus `core.report_brand.word_context`.

Run from the backend directory:  python scripts/goldline_build_word_templates.py
"""

import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))
from core import report_brand as brand  # noqa: E402

OUT_DIR = BACKEND / "core" / "templates" / "core"
LOGO = BACKEND / "core" / "templates" / "core" / "goldline-grc-mark.png"
FONT = "Calibri"

NAVY, BODY, TEXT, GOLD = brand.NAVY, brand.BODY, brand.TEXT, brand.GOLD
BORDER, CREAM, WHITE = brand.BORDER, brand.CREAM, brand.WHITE

STRINGS = {
    "en": {
        "eyebrow": "Compliance assessment",
        "subtitle": "Mapped to {{ audit.framework.name }}",
        "report_line": "Compliance assessment report",
        "company_assessed": "Company assessed",
        "prepared_for": "Prepared for",
        "engagement": "Engagement",
        "report_date": "Report date",
        "version": "Version",
        "assessor": "Assessor",
        "conf_lead": "Confidential working document.",
        "conf_body": " Prepared for {{ client }} and intended only for its management and auditors. "
        "Do not share it with third parties without written agreement from {{ brand.company }}. "
        "Governed by the law of {{ brand.jurisdiction }}.",
        "contents": "Contents",
        "toc_hint": "Right-click here and choose Update Field to refresh the contents.",
        "doc_control": "Document control",
        "date": "Date",
        "author": "Author",
        "change": "Change",
        "first_issue": "First issue",
        "page": "Page",
        "s_summary": "1. Summary",
        "scope": "Assessment scope:",
        "description": "Description:",
        "observations": "Observations:",
        "results": "Results overview",
        "req_total": "Number of requirements:",
        "ratings": ["Met", "Partially met", "Not met", "Not applicable", "Not assessed"],
        "assessor_comments": "Assessor comments:",
        "s_scopes": "2. Audited scopes",
        "scopes_intro": "The percentage shows how complete each section is, counting requirements rated Met or Not applicable.",
        "radar_intro": "The chart below shows the maturity the assessor found in each section.",
        "auditor_notes": "Assessor notes",
        "category": "Category",
        "gaps": "Gaps identified",
        "comments": "Comments",
        "s_controls": "3. Security controls",
        "controls_count": "Identified security controls:",
        "p1": "Main priorities (P1)",
        "p1_intro": "Security controls not yet in place that the assessor has made a priority:",
        "control": "Control",
        "status": "Status",
        "covered": "Requirements covered",
        "observations_h": "Assessor observations",
        "s_details": "4. Detailed results",
        "details_intro": "What this covers:",
        "details_body": " each assessable requirement with its rating, the observation recorded and the linked controls.",
        "iso_ref": "ISO reference",
        "req_evidence": "Requirement and evidence",
        "rating": "Rating",
        "observation": "Observation:",
        "applied": "Applied controls:",
        "s_plan": "5. Full action plan",
        "priority": "Priority",
        "due": "Due date",
        "s_appendix": "6. Appendix",
        "appendix": [
            "Details of the requirements and the evidence attached",
            "The remediation plan (list of controls)",
            "Assessor background and approach",
        ],
        "title": "Compliance assessment report",
    },
    "fr": {
        "eyebrow": "Évaluation de conformité",
        "subtitle": "Selon {{ audit.framework.name }}",
        "report_line": "Rapport d'évaluation de conformité",
        "company_assessed": "Entreprise évaluée",
        "prepared_for": "Préparé pour",
        "engagement": "Mission",
        "report_date": "Date du rapport",
        "version": "Version",
        "assessor": "Évaluateur",
        "conf_lead": "Document de travail confidentiel.",
        "conf_body": " Préparé pour {{ client }} et destiné uniquement à sa direction et à ses auditeurs. "
        "Ne pas le communiquer à des tiers sans l'accord écrit de {{ brand.company }}. "
        "Régi par le droit de l'Angleterre et du pays de Galles.",
        "contents": "Sommaire",
        "toc_hint": "Clic droit ici puis Mettre à jour les champs pour actualiser le sommaire.",
        "doc_control": "Suivi du document",
        "date": "Date",
        "author": "Auteur",
        "change": "Modification",
        "first_issue": "Première émission",
        "page": "Page",
        "s_summary": "1. Synthèse",
        "scope": "Périmètre évalué :",
        "description": "Description :",
        "observations": "Observations :",
        "results": "Vue d'ensemble des résultats",
        "req_total": "Nombre d'exigences :",
        "ratings": ["Satisfait", "Partiellement satisfait", "Non satisfait", "Non applicable", "Non évalué"],
        "assessor_comments": "Commentaires de l'évaluateur :",
        "s_scopes": "2. Périmètres audités",
        "scopes_intro": "Le pourcentage indique l'avancement de chaque section, en comptant les exigences satisfaites ou non applicables.",
        "radar_intro": "Le graphique ci-dessous montre la maturité constatée par l'évaluateur pour chaque section.",
        "auditor_notes": "Notes de l'évaluateur",
        "category": "Catégorie",
        "gaps": "Écarts identifiés",
        "comments": "Commentaires",
        "s_controls": "3. Mesures de sécurité",
        "controls_count": "Mesures de sécurité identifiées :",
        "p1": "Priorités principales (P1)",
        "p1_intro": "Mesures de sécurité non encore en place que l'évaluateur a classées prioritaires :",
        "control": "Mesure",
        "status": "Statut",
        "covered": "Exigences couvertes",
        "observations_h": "Observations de l'évaluateur",
        "s_details": "4. Résultats détaillés",
        "details_intro": "Contenu :",
        "details_body": " chaque exigence évaluable avec sa note, l'observation enregistrée et les mesures liées.",
        "iso_ref": "Référence ISO",
        "req_evidence": "Exigence et preuves",
        "rating": "Évaluation",
        "observation": "Observation :",
        "applied": "Mesures appliquées :",
        "s_plan": "5. Plan d'action complet",
        "priority": "Priorité",
        "due": "Échéance",
        "s_appendix": "6. Annexe",
        "appendix": [
            "Détail des exigences et des preuves jointes",
            "Le plan de remédiation (liste des mesures)",
            "Parcours et approche de l'évaluateur",
        ],
        "title": "Rapport d'évaluation de conformité",
    },
}

RATING_KEYS = ["compliant", "partially_compliant", "non_compliant", "not_applicable", "not_assessed"]
RATING_COUNTS = [
    "{{ req.compliant }}",
    "{{ req.partially_compliant }}",
    "{{ req.non_compliant }}",
    "{{ req.not_applicable }}",
    "{{ req.not_assessed }}",
]


# ------------------------------------------------------------------ helpers


def set_run_font(rpr, size=None, colour=None, bold=None, italic=None):
    """Explicit Calibri, size and colour, with theme overrides removed."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        del fonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)
    if colour is not None:
        for old in rpr.findall(qn("w:color")):
            rpr.remove(old)
        c = OxmlElement("w:color")
        c.set(qn("w:val"), colour)
        rpr.append(c)
    if size is not None:
        for tag in ("w:sz", "w:szCs"):
            for old in rpr.findall(qn(tag)):
                rpr.remove(old)
            el = OxmlElement(tag)
            el.set(qn("w:val"), str(int(round(size * 2))))
            rpr.append(el)
    for tag, value in (("w:b", bold), ("w:i", italic)):
        if value is None:
            continue
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
        el = OxmlElement(tag)
        if not value:
            el.set(qn("w:val"), "0")
        rpr.append(el)


def style_font(style, size, colour, bold=False):
    rpr = style.element.get_or_add_rPr()
    set_run_font(rpr, size=size, colour=colour, bold=bold, italic=False)


def run(paragraph, text, size=10.5, colour=BODY, bold=False, italic=False):
    r = paragraph.add_run(text)
    set_run_font(r._r.get_or_add_rPr(), size=size, colour=colour, bold=bold, italic=italic)
    return r


def spacing(paragraph, before=0, after=0, line=None):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def bottom_rule(paragraph, size_pt, colour=GOLD, space=1):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(int(size_pt * 8)))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), colour)
    bdr.append(bottom)
    ppr.append(bdr)


def field(paragraph, code, size, colour, placeholder=""):
    def fld(kind):
        r = paragraph.add_run()
        set_run_font(r._r.get_or_add_rPr(), size=size, colour=colour)
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._r.append(el)

    fld("begin")
    r = paragraph.add_run()
    set_run_font(r._r.get_or_add_rPr(), size=size, colour=colour)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {code} "
    r._r.append(instr)
    fld("separate")
    run(paragraph, placeholder, size=size, colour=colour)
    fld("end")


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table_borders_and_padding(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")  # eighths of a point: 0.5 pt
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), "80")  # 4 pt
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblpr.append(width)


def cell_text(cell, text, size=9, colour=TEXT, bold=False, align=None):
    p = cell.paragraphs[0]
    spacing(p, 0, 0, 1.0)
    if align is not None:
        p.alignment = align
    run(p, text, size=size, colour=colour, bold=bold)
    return p


def new_table(doc, columns, widths_mm=None):
    t = doc.add_table(rows=0, cols=columns)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    table_borders_and_padding(t)
    if widths_mm:
        for i, w in enumerate(widths_mm):
            t.columns[i].width = Mm(w)
    return t


def header_row(table, labels, widths_mm=None):
    row = table.add_row()
    trpr = row._tr.get_or_add_trPr()
    rep = OxmlElement("w:tblHeader")  # repeat on every page
    trpr.append(rep)
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)
    for i, label in enumerate(labels):
        shade(row.cells[i], NAVY)
        cell_text(row.cells[i], label, colour=WHITE, bold=True)
        if widths_mm:
            row.cells[i].width = Mm(widths_mm[i])
    return row


def body_row(table, values, widths_mm=None, fills=None, bold_first=False):
    row = table.add_row()
    for i, value in enumerate(values):
        if fills and fills[i]:
            shade(row.cells[i], fills[i])
        cell_text(row.cells[i], value, bold=bold_first and i == 0)
        if widths_mm:
            row.cells[i].width = Mm(widths_mm[i])
    return row


def label_para(doc, lead, body, size=10.5):
    p = doc.add_paragraph()
    run(p, lead, size=size, colour=TEXT, bold=True)
    run(p, " ", size=size, colour=BODY)
    run(p, body, size=size, colour=BODY)  # own run: docxtpl {{r }} tags need it
    return p


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    run(p, text, size=15 if level == 1 else 12.5, colour=NAVY, bold=True)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    spacing(p)
    p.add_run().add_break(WD_BREAK.PAGE)


def image_para(doc, tag):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "{{ " + tag + " }}")



# Word rejects formatting blocks whose children are out of schema order, so
# everything appended above is put back in order before saving.
SCHEMA_ORDER = {
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "trPr": "cnfStyle divId gridBefore gridAfter wBefore wAfter cantSplit trHeight tblHeader tblCellSpacing jc hidden",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
}


def reorder(root):
    for parent_tag, order in SCHEMA_ORDER.items():
        rank = {qn("w:" + name): i for i, name in enumerate(order.split())}
        for parent in root.iter(qn("w:" + parent_tag)):
            kids = list(parent)
            # Keep only the last of any duplicated property.
            seen = {}
            for k in kids:
                seen[k.tag] = k
            kids = [k for k in kids if seen.get(k.tag) is k]
            for k in list(parent):
                parent.remove(k)
            for k in sorted(kids, key=lambda k: rank.get(k.tag, len(rank))):
                parent.append(k)


def normalise(doc):
    reorder(doc.element)
    reorder(doc.styles.element)
    for sec in doc.sections:
        reorder(sec.header._element)
        reorder(sec.footer._element)

# ------------------------------------------------------------------- builder


def build(lang: str) -> Document:
    s = STRINGS[lang]
    doc = Document()

    # Page: A4, 20 mm side margins.
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(24), Mm(22)
    sec.header_distance, sec.footer_distance = Mm(10), Mm(10)

    # Styles.
    normal = doc.styles["Normal"]
    style_font(normal, 10.5, BODY)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)
    for level, size, before, after in ((1, 15, 18, 8), (2, 12.5, 12, 6)):
        st = doc.styles[f"Heading {level}"]
        style_font(st, size, NAVY, bold=True)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.keep_with_next = True
    for level in (1, 2):
        name = f"TOC {level}"
        st = doc.styles[name] if name in [x.name for x in doc.styles] else doc.styles.add_style(name, 1)
        st.base_style = normal
        style_font(st, 10.5, TEXT, bold=(level == 1))
        st.paragraph_format.left_indent = Mm(0 if level == 1 else 5)
        st.paragraph_format.space_after = Pt(2)
        tabs = st.paragraph_format.tab_stops
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER

        tabs.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    for name in ("List Bullet",):
        style_font(doc.styles[name], 10.5, BODY)

    # Running header: "Client  |  Report title", 7.5 pt grey, gold rule.
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    spacing(hp, 0, 0, 1.0)
    run(hp, "{{ client }}  |  {{ audit.name }}", size=7.5, colour=BODY)
    bottom_rule(hp, 0.75, GOLD, space=4)

    # Footer: centred company line, page number in 9.5 pt dark.
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(fp, 0, 0, 1.0)
    run(fp, "{{ brand.footer }}  " + s["page"] + " ", size=7.5, colour=BODY)
    field(fp, "PAGE", 9.5, TEXT, "1")

    # ------------------------------------------------------------- cover
    p = doc.add_paragraph()
    spacing(p, 0, 0, 1.0)
    p.add_run().add_picture(str(LOGO), width=Mm(20))
    p = doc.add_paragraph()
    spacing(p, 46, 6)
    run(p, s["eyebrow"].upper(), size=10, colour=GOLD, bold=True)
    p = doc.add_paragraph()
    spacing(p, 0, 2, 1.1)
    run(p, "{{ audit.name }}", size=22, colour=NAVY, bold=True)
    p = doc.add_paragraph()
    spacing(p, 0, 6, 1.1)
    run(p, s["subtitle"], size=15, colour=NAVY)
    p = doc.add_paragraph()
    spacing(p, 0, 6, 1.1)
    run(p, s["report_line"], size=12.5, colour=TEXT, bold=True)
    p = doc.add_paragraph()
    spacing(p, 0, 14, 1.0)
    bottom_rule(p, 2, GOLD, space=1)

    widths = (62, 108)
    t = new_table(doc, 2, widths)
    for label, key in (
        (s["company_assessed"], "company_assessed"),
        (s["prepared_for"], "prepared_for"),
        (s["engagement"], "engagement"),
        (s["report_date"], "date"),
        (s["version"], "version"),
        (s["assessor"], "assessor"),
    ):
        row = t.add_row()
        shade(row.cells[0], CREAM)
        cell_text(row.cells[0], label, bold=True)
        cell_text(row.cells[1], "{{r fields." + key + " }}")
        row.cells[0].width, row.cells[1].width = Mm(widths[0]), Mm(widths[1])
    p = doc.add_paragraph()
    spacing(p, 14, 0, 1.3)
    run(p, s["conf_lead"], size=9, colour=TEXT, bold=True)
    run(p, s["conf_body"], size=9, colour=BODY)
    page_break(doc)

    # ---------------------------------------------------------- contents
    p = doc.add_paragraph()
    spacing(p, 0, 10, 1.1)
    run(p, s["contents"], size=15, colour=NAVY, bold=True)
    p = doc.add_paragraph()
    field(p, 'TOC \\o "1-2" \\h \\z \\u', 10.5, TEXT, s["toc_hint"])
    page_break(doc)

    # -------------------------------------------------- document control
    p = doc.add_paragraph()
    spacing(p, 0, 10, 1.1)
    run(p, s["doc_control"], size=15, colour=NAVY, bold=True)
    widths = (25, 30, 70, 45)
    t = new_table(doc, 4, widths)
    header_row(t, (s["version"], s["date"], s["author"], s["change"]), widths)
    row = t.add_row()
    for i, value in enumerate(
        ("{{r fields.version }}", "{{r fields.date }}", "{{r fields.assessor }}", s["first_issue"])
    ):
        cell_text(row.cells[i], value)
        row.cells[i].width = Mm(widths[i])
    page_break(doc)

    # ------------------------------------------------------------ summary
    heading(doc, s["s_summary"], 1)
    label_para(doc, s["scope"], "{{r fields.scope }}")
    label_para(doc, s["description"], "{{r fields.description }}")
    label_para(doc, s["observations"], "{{r fields.observation }}")
    heading(doc, s["results"], 2)
    label_para(doc, s["req_total"], "{{ req.total }}")
    widths = (34,) * 5
    t = new_table(doc, 5, widths)
    header_row(t, s["ratings"], widths)
    row = t.add_row()
    for i, (key, value) in enumerate(zip(RATING_KEYS, RATING_COUNTS)):
        shade(row.cells[i], brand.RATING_FILL[key])
        cell_text(row.cells[i], value, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        row.cells[i].width = Mm(widths[i])
    image_para(doc, "compliance_donut")
    p = doc.add_paragraph()
    run(p, s["assessor_comments"], colour=TEXT, bold=True)
    run(p, " ")
    run(p, "{{r comments }}")

    # ------------------------------------------------------------- scopes
    heading(doc, s["s_scopes"], 1)
    p = doc.add_paragraph()
    run(p, s["scopes_intro"])
    image_para(doc, "completion_bar")
    p = doc.add_paragraph()
    run(p, s["radar_intro"])
    image_para(doc, "category_radar")
    heading(doc, s["auditor_notes"], 2)
    widths = (70, 30, 70)
    t = new_table(doc, 3, widths)
    header_row(t, (s["category"], s["gaps"], s["comments"]), widths)
    r = t.add_row()
    cell_text(r.cells[0], "{%tr for item in drifts_per_domain %}")
    body_row(t, ("{{ item.name }}", "{{ item.drift_count }}", "{{r comments }}"), widths)
    r = t.add_row()
    cell_text(r.cells[0], "{%tr endfor %}")

    # ----------------------------------------------------------- controls
    heading(doc, s["s_controls"], 1)
    label_para(doc, s["controls_count"], "{{ ac_count }}")
    image_para(doc, "chart_controls")
    heading(doc, s["p1"], 2)
    p = doc.add_paragraph()
    run(p, s["p1_intro"])
    widths = (38, 54, 24, 24, 30)
    t = new_table(doc, 5, widths)
    header_row(t, (s["control"], s["description"].rstrip(" :"), s["category"], s["status"], s["covered"]), widths)
    r = t.add_row()
    cell_text(r.cells[0], "{%tr for item in p1_controls %}")
    body_row(
        t,
        ("{{ item.name }}", "{{ item.description }}", "{{ item.category }}", "{{ item.status }}", "{{ item.coverage }}"),
        widths,
        bold_first=True,
    )
    r = t.add_row()
    cell_text(r.cells[0], "{%tr endfor %}")
    heading(doc, s["observations_h"], 2)
    p = doc.add_paragraph()
    run(p, "{{r comments }}")

    # ---------------------------------------------------- detailed results
    heading(doc, s["s_details"], 1)
    label_para(doc, s["details_intro"], s["details_body"].strip())
    widths = (24, 112, 34)
    t = new_table(doc, 3, widths)
    header_row(t, (s["iso_ref"], s["req_evidence"], s["rating"]), widths)
    r = t.add_row()
    cell_text(r.cells[0], "{%tr for ra in requirement_assessments %}")
    row = t.add_row()
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))
    cell_text(row.cells[0], "{{ ra.ref_id }}")
    c = row.cells[1]
    cell_text(c, "{{ ra.name }}", colour=TEXT, bold=True)
    for lead, value in ((s["observation"], "ra.observation"), (s["applied"], "ra.applied_controls")):
        p1 = c.add_paragraph()
        spacing(p1, 0, 0, 1.0)
        run(p1, "{%p if " + value + ' != "-" %}', size=9)
        p2 = c.add_paragraph()
        spacing(p2, 2, 0, 1.0)
        run(p2, lead, size=9, colour=TEXT, bold=True)
        run(p2, " {{ " + value + " }}", size=9, colour=BODY)
        p3 = c.add_paragraph()
        spacing(p3, 0, 0, 1.0)
        run(p3, "{%p endif %}", size=9)
    cell_text(row.cells[2], "{% cellbg ra.rating_fill %}{{ ra.rating_label }}")
    for i in range(3):
        row.cells[i].width = Mm(widths[i])
    r = t.add_row()
    cell_text(r.cells[0], "{%tr endfor %}")

    # --------------------------------------------------------- action plan
    heading(doc, s["s_plan"], 1)
    widths = (32, 40, 15, 20, 20, 18, 25)
    t = new_table(doc, 7, widths)
    header_row(
        t,
        (s["control"], s["description"].rstrip(" :"), s["priority"], s["category"], s["status"], s["due"], s["covered"]),
        widths,
    )
    r = t.add_row()
    cell_text(r.cells[0], "{%tr for item in full_controls %}")
    body_row(
        t,
        (
            "{{ item.name }}",
            "{{ item.description }}",
            "{{ item.prio }}",
            "{{ item.category }}",
            "{{ item.status }}",
            "{{ item.eta }}",
            "{{ item.coverage }}",
        ),
        widths,
        bold_first=True,
    )
    r = t.add_row()
    cell_text(r.cells[0], "{%tr endfor %}")

    # ------------------------------------------------------------ appendix
    heading(doc, s["s_appendix"], 1)
    for item in s["appendix"]:
        p = doc.add_paragraph(style="List Bullet")
        run(p, item)

    # Metadata and "update the contents on open".
    cp = doc.core_properties
    cp.author = cp.last_modified_by = brand.COMPANY
    cp.title = s["title"]
    cp.comments = ""
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    # updateFields sits before these in the settings sequence.
    later = [c for c in settings if c.tag.split("}")[1] in (
        "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
        "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
        "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
        "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
        "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")]
    if later:
        later[0].addprevious(upd)
    else:
        settings.append(upd)
    normalise(doc)
    return doc


def set_company(path: Path):
    """python-docx leaves app.xml alone; set Company there too."""
    with zipfile.ZipFile(path) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in entries:
            if info.filename == "docProps/app.xml":
                xml = data.decode()
                xml = re.sub(r"<Company>.*?</Company>|<Company/>", "", xml)
                xml = xml.replace("</Properties>", f"<Company>{brand.COMPANY}</Company></Properties>")
                data = xml.encode()
            z.writestr(info, data)


def main():
    for lang, names in (("en", ("audit_report_template_en.docx", "audit_report_template.docx")), ("fr", ("audit_report_template_fr.docx",))):
        for name in names:
            out = OUT_DIR / name
            build(lang).save(out)
            set_company(out)
            print("wrote", out)


if __name__ == "__main__":
    main()
